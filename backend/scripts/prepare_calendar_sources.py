"""Prepare auditable calendar imports; never writes to the app database.

Reader dependencies: pymupdf, openpyxl. Output is private runtime data and
must stay outside source control. Unknown/ambiguous dates stay unresolved.
"""
import argparse
from collections import Counter
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import unicodedata
from zoneinfo import ZoneInfo


MONTHS = {name: i for i, name in enumerate(
    ['janvier', 'fevrier', 'mars', 'avril', 'mai', 'juin', 'juillet',
     'aout', 'septembre', 'octobre', 'novembre', 'decembre'], 1)}


def fold(value):
    return ''.join(c for c in unicodedata.normalize('NFD', value.lower())
                   if unicodedata.category(c) != 'Mn')


def parse_dates(value):
    """Return inclusive dates only for explicit, unambiguous source values."""
    if isinstance(value, (datetime, date)):
        d = value.date() if isinstance(value, datetime) else value
        return d.isoformat(), d.isoformat(), None
    text = fold(str(value or '')).strip()
    if re.fullmatch(r'\d{4}', text):
        return None, None, 'Year only; no event dates'
    if re.search(r'\d{5,}', text):
        return None, None, 'Malformed year; review original source'
    year = re.search(r'\b(20\d{2})\b', text)
    month = next((v for k, v in MONTHS.items() if re.search(r'\b'+k+r'\b', text)), None)
    if not year or not month:
        return None, None, 'Missing year or ambiguous date; human review required'
    prefix = text[:next(re.search(r'\b'+k+r'\b', text).start()
                        for k, v in MONTHS.items() if v == month)]
    days = [int(x) for x in re.findall(r'\d+', prefix)]
    # Different named months or parenthetical optional dates need review.
    if len(days) not in (1, 2, 3) or '(' in prefix or sum(k in text for k in MONTHS) != 1:
        return None, None, 'Complex date range; human review required'
    try:
        start, end = date(int(year[1]), month, days[0]), date(int(year[1]), month, days[-1])
        if end < start:
            raise ValueError('Reversed date range')
        return start.isoformat(), end.isoformat(), None
    except ValueError as exc:
        return None, None, str(exc)


def snapshot(path, target, name):
    copied = target / 'sources' / name
    copied.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, copied)
    return {'file': str(copied), 'sha256': hashlib.sha256(copied.read_bytes()).hexdigest()}


def schedule(path, group, source):
    import pymupdf
    result = []
    weekdays = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche']
    pattern = re.compile(r'(\d{4}-\d{2}-\d{2})\s+(\w+)\s+(\d{2}:\d{2})\s+(\d{2}:\d{2})\s+(INSEP|SNMT)')
    with pymupdf.open(path) as doc:
        for page_number, page in enumerate(doc, 1):
            text = page.get_text(sort=True)
            matches = list(pattern.finditer(text))
            if len(re.findall(r'\d{4}-\d{2}-\d{2}', text)) != len(matches):
                raise ValueError(f'Unparsed schedule rows in {group}, page {page_number}')
            for row, match in enumerate(matches, 1):
                day, weekday, start, end, venue = match.groups()
                d = date.fromisoformat(day)
                if weekdays[d.weekday()] != weekday or end <= start:
                    raise ValueError(f'Invalid schedule row: {match[0]}')
                local_start = datetime.fromisoformat(day+'T'+start).replace(tzinfo=ZoneInfo('Europe/Paris'))
                local_end = datetime.fromisoformat(day+'T'+end).replace(tzinfo=ZoneInfo('Europe/Paris'))
                result.append({'source_key': f'kcpm:2026-2027:{group}:{day}',
                    'club_key': 'kcpm', 'group_key': group, 'kind': 'training',
                    'local_date': day, 'start_local': start, 'end_local': end,
                    'timezone': 'Europe/Paris', 'start_at': local_start.astimezone(timezone.utc).isoformat(),
                    'end_at': local_end.astimezone(timezone.utc).isoformat(), 'venue_code': venue,
                    'source': {**source, 'page': page_number, 'row': row},
                    'publication_status': 'staged', 'membership_assignments': []})
    if not result or len({r['source_key'] for r in result}) != len(result):
        raise ValueError(f'Empty schedule or duplicate group/date in {group}')
    return result


def workbook(path, source):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    raw, meets, resources, reviews = [], [], [], []
    for sheet_index, sheet in enumerate(wb):
        for row in sheet:
            cells = [{'cell': c.coordinate, 'value': c.value.isoformat() if isinstance(c.value, (datetime, date)) else c.value,
                      'url': c.hyperlink.target if c.hyperlink else None}
                     for c in row if c.value is not None or c.hyperlink]
            if not cells:
                continue
            ref = {**source, 'sheet': sheet.title, 'row': row[0].row}
            raw.append({'source': ref, 'cells': cells})
            if sheet_index == 0 and row[0].row > 3 and sheet.cell(row[0].row, 3).value:
                n = row[0].row
                start, end, issue = parse_dates(sheet.cell(n, 1).value)
                links = []
                for cell in cells:
                    if cell['url']:
                        links.append({**cell, 'visibility': 'club_only' if 'chat.whatsapp.com' in cell['url'] else 'review_required'})
                meet = {'snapshot_row_key': f'workbook:competitions:{n}',
                        'title': str(sheet.cell(n, 3).value).strip(),
                        'location_raw': sheet.cell(n, 2).value,
                        'level_raw': sheet.cell(n, 4).value,
                        'date_raw': next((c['value'] for c in cells if c['cell'] == f'A{n}'), None),
                        'start_date': start, 'end_date_inclusive': end,
                        'end_date_exclusive': (date.fromisoformat(end)+timedelta(days=1)).isoformat() if end else None,
                        'registration_deadline_raw': next((c['value'] for c in cells if c['cell'] == f'F{n}'), None),
                        'notes': sheet.cell(n, 7).value, 'links': links, 'source': ref,
                        'publication_status': 'needs_review' if issue else 'staged_unverified',
                        'date_issue': issue}
                meets.append(meet)
                if issue:
                    reviews.append({'source': ref, 'title': meet['title'], 'issue': issue})
            elif sheet_index == 2 and row[0].row > 1:
                n = row[0].row
                resources.append({'type': sheet.cell(n, 1).value, 'title': sheet.cell(n, 2).value,
                                  'url': sheet.cell(n, 3).hyperlink.target if sheet.cell(n, 3).hyperlink else None,
                                  'notes': sheet.cell(n, 4).value, 'source': ref,
                                  'verification_status': 'unverified'})
    wb.close()
    return raw, meets, resources, reviews


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--beginners', type=Path, required=True)
    p.add_argument('--advanced', type=Path, required=True)
    p.add_argument('--workbook', type=Path, required=True)
    p.add_argument('--out', type=Path, required=True)
    a = p.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    sources = {key: snapshot(path, a.out, name) for key, path, name in [
        ('beginners', a.beginners, 'beginners-2026-2027.pdf'),
        ('advanced', a.advanced, 'advanced-2026-2027.pdf'),
        ('workbook', a.workbook, 'competitions.xlsx')]}
    sessions = schedule(a.beginners, 'beginners', sources['beginners']) + schedule(a.advanced, 'advanced', sources['advanced'])
    raw, meets, resources, reviews = workbook(a.workbook, sources['workbook'])
    payload = {'format_version': 1, 'state': 'prepared_not_loaded_into_app',
               'sources': sources, 'training_occurrences': sessions,
               'competition_candidates': meets, 'resources': resources,
               'workbook_rows': raw, 'review_queue': reviews,
               'refresh': {'mode': 'manual_snapshot', 'live_source_url': None}}
    (a.out / 'import-bundle.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'groups': dict(Counter(s['group_key'] for s in sessions)),
        'venues': dict(Counter(s['venue_code'] for s in sessions)),
        'competition_candidates': len(meets), 'resources': len(resources),
        'date_review_count': len(reviews), 'state': payload['state']}))


if __name__ == '__main__':
    main()
